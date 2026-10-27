"""Generate a fresh report while retaining the original Word cover's XML."""
import copy
import io
import os
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from .model import ROOT, safe_name, selected_nodes, validate_project

TEMPLATE = ROOT / 'A-B-Průvodní list a souhrnná zpráva ZMĚNA.docx'


def set_text(paragraph, text):
    """Keep paragraph, run formatting and floating drawings; replace text only."""
    texts = paragraph._p.findall('.//' + qn('w:t'))
    if texts:
        texts[0].text = text
        texts[0].set(qn('xml:space'), 'preserve')
        for item in texts[1:]:
            item.text = ''
    else:
        paragraph.add_run(text)


def generate(project):
    project = validate_project(project, require_name=True)
    nodes = selected_nodes(project['selected'])
    if not nodes:
        raise ValueError('Vyberte alespoň jednu kapitolu.')
    doc = Document(TEMPLATE)
    paragraphs = doc.paragraphs
    set_text(paragraphs[3], project['name'])
    set_text(paragraphs[4], f"{project['stage']}\t z.č. {project['number']}")
    # Keep the cover's anchored image and all paragraph/run geometry.
    set_text(paragraphs[6], 'Průvodní list a souhrnná technická zpráva')
    set_text(paragraphs[16], f"Datum: {project['date']}\t\tVypracoval: {project['author']}")
    if project['cover_image']:
        image = Path(project['cover_image']).expanduser()
        if not image.is_file():
            raise ValueError('Vybraný obrázek není dostupný. Vyberte jej znovu.')
        drawing = paragraphs[6]._p.find('.//' + qn('a:blip'))
        if drawing is None:
            raise ValueError('V předloze chybí obrázek titulní strany.')
        from PIL import Image
        with Image.open(image) as im:
            if im.width * im.height > 40_000_000:
                raise ValueError('Obrázek je příliš velký (maximum 40 megapixelů).')
            # Fit image into original frame, preserving aspect ratio.
            from PIL import ImageOps
            frame = ImageOps.contain(im.convert('RGB'), (1600, 1000))
            canvas = Image.new('RGB', (1600, 1000), 'white')
            canvas.paste(frame, ((1600-frame.width)//2, (1000-frame.height)//2))
            stream = io.BytesIO()
            canvas.save(stream, format='PNG')
        stream.seek(0)
        rid, _ = doc.part.get_or_add_image(stream)
        drawing.set(qn('r:embed'), rid)

    # The first page consists of paragraphs 0..16; body starts with a page break.
    body = doc._element.body
    for element in list(body)[17:]:
        if element.tag != qn('w:sectPr'):
            body.remove(element)
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            set_text(paragraph, f"{project['name']}\tPrůvodní list a souhrnná zpráva")
    doc.core_properties.title = project['name']
    doc.core_properties.subject = 'Připravená osnova technické zprávy'
    doc.core_properties.author = project['author']
    doc.core_properties.last_modified_by = project['author']
    doc.core_properties.comments = ''
    for level in range(1, 6):
        style = doc.styles[f'Heading {level}']
        style.font.name = 'Calibri'
        style.font.size = Pt(max(11, 17 - level))
        style.paragraph_format.keep_with_next = True

    fields = {'A1.1': [('Název stavby', project['name']), ('Místo stavby', project['location'])],
              'A1.2': [('Stavebník / investor', project['investor'])],
              'A1.3': [('Zpracovatel', project['author'])]}

    def append(items, level=1):
        for node in items:
            p = doc.add_heading(f"{node['code']} {node['title']}", level=min(level, 5))
            if level == 1:
                p.paragraph_format.page_break_before = True
            if node['explicit']:
                for label, value in fields.get(node['id'], []):
                    doc.add_paragraph(f"{label}: {value or '[DOPLNIT]'}")
                if not node['children']:
                    doc.add_paragraph('[DOPLNIT: ' + node.get('hint', node['title']).strip(' .') + '.]')
            append(node['children'], level + 1)
    append(nodes)
    # Remove orphaned image relationships from deleted filled-in report pages.
    used = {el.get(qn('r:embed')) for el in doc._element.iter() if el.get(qn('r:embed'))}
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/image') and rid not in used:
            doc.part.drop_rel(rid)
    result = io.BytesIO()
    doc.save(result)
    return result.getvalue()


def write_version(content, directory, filename):
    folder = Path(directory).expanduser().resolve()
    if not folder.is_dir():
        raise ValueError('Výstupní složka neexistuje. Vyberte existující složku.')
    stem = safe_name(Path(filename).stem)
    for version in range(1, 10000):
        path = folder / (stem + (f'_{version:02d}' if version > 1 else '') + '.docx')
        try:
            with path.open('xb') as f:
                f.write(content)
            return path
        except FileExistsError:
            continue
    raise ValueError('Pro tento název už existuje příliš mnoho verzí.')
