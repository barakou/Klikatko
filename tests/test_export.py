import io
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from klikatko.export import TEMPLATE, cover_info, generate, template_cover_image, write_version


class ExportTests(unittest.TestCase):
    def project(self, **overrides):
        project = {
            'name': 'Demo Park Renewal',
            'number': 'JOB-2048',
            'location': 'Northfield, Plot 17',
            'investor': 'Example Client',
            'author': 'Casey Example',
            'date': '2026-10-27',
            'stage': 'Permit documentation',
            'selected': ['B3.4.SO101.b'],
        }
        project.update(overrides)
        return project

    def test_deep_selection_keeps_ancestors_and_excludes_siblings(self):
        output = Document(io.BytesIO(generate(self.project())))
        headings = [p.text for p in output.paragraphs if p.style.name.startswith('Heading')]
        self.assertIn('B Souhrnná technická zpráva', headings)
        self.assertIn('B3 Základní stavebně technické a technologické řešení', headings)
        self.assertIn('B3.4 Základní technický popis stavebních objektů', headings)
        self.assertIn('SO 101 Komunikace', headings)
        self.assertIn('b) Popis navrženého stavebně konstrukčního a technického řešení', headings)
        self.assertFalse(any('SO 401' in heading for heading in headings))
        self.assertFalse(any('a) Popis stávajícího stavu' in heading for heading in headings))
        self.assertFalse(any('A4.' in heading for heading in headings))
        self.assertFalse(any('A4.' in p.text for p in output.paragraphs))

    def test_filled_body_is_removed_and_cover_uses_project_values(self):
        output = Document(io.BytesIO(generate(self.project())))
        text = '\n'.join(p.text for p in output.paragraphs)
        self.assertIn('Demo Park Renewal', text)
        self.assertIn('JOB-2048', text)
        self.assertIn('Permit documentation', text)
        self.assertIn('2026-10-27', text)
        self.assertIn('Casey Example', text)
        self.assertNotIn('Jedná se o II. Etapu rekonstrukce ulice Dlouhé', text)
        self.assertNotIn('Ve stavbě nebyly zjištěny žádné památné stromy', text)

    def test_cover_anchor_and_original_margins_are_preserved(self):
        source = Document(TEMPLATE)
        output = Document(io.BytesIO(generate(self.project())))
        self.assertIsNotNone(source.paragraphs[6]._p.find('.//' + qn('wp:anchor')))
        self.assertIsNotNone(output.paragraphs[6]._p.find('.//' + qn('wp:anchor')))
        source_margins = [
            (s.page_width, s.page_height, s.top_margin, s.bottom_margin, s.left_margin, s.right_margin)
            for s in source.sections
        ]
        output_margins = [
            (s.page_width, s.page_height, s.top_margin, s.bottom_margin, s.left_margin, s.right_margin)
            for s in output.sections
        ]
        self.assertEqual(output_margins, source_margins)

    def test_write_version_sanitizes_filename_and_never_overwrites(self):
        with tempfile.TemporaryDirectory() as directory:
            first = write_version(b'first', directory, 'unsafe?name.docx')
            second = write_version(b'second', directory, 'unsafe?name.docx')
            self.assertEqual(first.name, 'unsafe-name.docx')
            self.assertEqual(second.name, 'unsafe-name_02.docx')
            self.assertEqual(first.read_bytes(), b'first')
            self.assertEqual(second.read_bytes(), b'second')
            self.assertEqual(first.parent, Path(directory).resolve())

    def test_write_version_keeps_dots_inside_the_name(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_version(b'x', directory, 'A-B Rekonstrukce Dlouhé ulice, II. etapa')
            self.assertEqual(path.name, 'A-B Rekonstrukce Dlouhé ulice, II. etapa.docx')

    def test_cover_info_reads_letterhead_from_template(self):
        info = cover_info()
        self.assertTrue(info['letterhead'][0][0]['text'].startswith('ING. JIŘÍ KOUDELKA'))
        self.assertEqual(info['letterhead'][0][0]['size'], 24)
        self.assertTrue(any(run['link'] for run in info['letterhead'][1]))
        self.assertEqual(info['attachment'], 'Příloha A-B')
        self.assertEqual(template_cover_image()[1:4], b'PNG')


if __name__ == '__main__':
    unittest.main()
